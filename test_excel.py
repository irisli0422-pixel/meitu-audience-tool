#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_excel.py —— 本地自测脚本（不调用 AI，不消耗接口次数）

作用：用一份"假的 AI 返回数据"跑一遍 Excel 生成逻辑，
      验证生成出来的 .xlsx 结构、合并单元格、配色是否符合 old_skill 标准格式。

用法：.venv/bin/python test_excel.py
"""

import os

import openpyxl

import excel_builder
import history_store
import skill_rules

# ============================================================
# 模拟一份 AI 返回的人群包数据（结构与 skill_rules.OUTPUT_SCHEMA 一致）
# ============================================================
MOCK_DATA = {
    "brand_name": "测试品牌 TestBrand",
    "brand_name_en": "TESTBRAND",
    "colors": {"title": "#1A1A1A", "section": "2E5C8A", "header": "3B7DD8", "basic": "E8F0FE"},
    "data_notes": [
        "1. 整体逻辑通过高端护肤消费人群特征匹配美图人群画像；",
        "2. 涉及的APP：美图系场景（美图秀秀、美颜相机等）；",
        "3. 人群包采用近6个月中符合圈选规则的人群；",
        "4. 品牌背景信息：测试品牌主打抗老精华，定位高端。",
        "5. 本次核心场景：晨晚间护肤、妆前打底、医美术后修护。",
    ],
    "summary_text": "【基础定向 ∪ 影像识别护肤场景人群】∩【本竞品人群包 ∪ 护肤进阶人群包 ∪ 种子人群包】",
    "basic_targeting": [
        {"dimension": "性别维度", "logic": "女性为主（占比85%），整体男女全覆盖", "scale": ""},
        {"dimension": "年龄维度", "logic": "25-35岁为核心人群，兼顾22-45岁", "scale": ""},
        {"dimension": "兴趣爱好维度", "logic": "安装应用偏好：美妆电商、医美平台\n社区兴趣偏好：护肤测评、成分党", "scale": ""},
        {"dimension": "中高消费维度", "logic": "中高消费力（月收入15000+），对高端护肤有消费意愿", "scale": ""},
    ],
    "segments": [
        {
            "name": "本竞品人群包",
            "color": "FFF3E0",
            "tags": [
                {"tag": "广告点击行为-品牌", "logic": "在美图系平台的广告点击行为数据判断出用户对相关广告感兴趣，广告点击兴趣=雅诗兰黛/兰蔻/赫莲娜，频次=180天", "scale": "320w"},
                {"tag": "影像识别-竞品品牌", "logic": "近180天内通过影像识别侦察到【雅诗兰黛、兰蔻、赫莲娜、珀莱雅、薇诺娜】品牌形象的用户", "scale": "480w"},
                {"tag": "LBS-专柜到店", "logic": "近180天内通过LBS及影像识别侦察到【高端百货美妆专柜、丝芙兰门店、调色师门店】的用户，频次≥2次", "scale": "260w"},
                {"tag": "社区互动-品牌话题", "logic": "近180天内在美图社区中发布或互动【抗老精华、成分党测评、护肤空瓶记】相关话题及图片的用户", "scale": "210w"},
            ],
        },
        {
            "name": "护肤进阶人群包",
            "color": "E8F5E9",
            "tags": [
                {"tag": "影像识别-护肤场景", "logic": "影像识别用户近180天内图片中高频出现护肤场景：[梳妆台、精华瓶、面膜敷脸、化妆棉、护肤流程记录]", "scale": "650w"},
                {"tag": "原生需求-美肤功能", "logic": "近180天内在美图秀秀产品产生原生需求【磨皮、祛斑祛痘、素颜美肌、水光肌】配方/滤镜/贴纸的用户", "scale": "880w"},
                {"tag": "广告点击-护肤兴趣", "logic": "在美图系平台的广告点击行为数据判断出用户对相关广告感兴趣，广告点击兴趣=抗老/紧致/淡纹，频次=90天", "scale": "420w"},
                {"tag": "社区互动-护肤日常", "logic": "近180天内在美图社区中发布或互动【护肤日常、早C晚A、医美恢复期】相关话题及图片的用户", "scale": "330w"},
            ],
        },
        {
            "name": "种子人群包",
            "color": "EDEDED",
            "tags": [
                {"tag": "历史互动人群", "logic": "2024-2026年度在美图秀秀APP点击过测试品牌广告、参与过品牌话题互动、使用过联名配方/贴纸/滤镜的人群（去重种子人群包）", "scale": "78w"},
                {"tag": "品牌搜索人群", "logic": "近365天内在美图系平台搜索过测试品牌/抗老精华/紧致面霜等关键词的用户", "scale": "54w"},
            ],
        },
    ],
    "remarks": [
        "1. 以上人群包量级为预估量级，实际投放时以系统最终圈选结果为准；",
        "2. 各人群包之间采用并集（∪）逻辑组合，基础定向与定制人群包之间采用交集（∩）逻辑；",
        "3. 影像识别基于美图AI图像识别技术，识别用户上传/编辑图片中的品牌LOGO、场景、物品等元素；",
        "4. LBS场景识别基于用户打开APP时的地理位置信息，匹配线下消费场所POI数据；",
        "5. 原生需求指用户在美图秀秀产品中主动使用的功能/配方/滤镜等行为数据；",
        "6. 品牌核心信息：主打抗老精华，高端定位，核心卖点为专利成分；",
        "7. 核心人群优先级：本竞品人群包(高转化) → 护肤进阶人群包(高复购) → 种子人群包(高传播)。",
    ],
    "competitors": ["雅诗兰黛", "兰蔻", "赫莲娜", "SK-II", "珀莱雅", "薇诺娜", "夸迪", "可复美"],
    "industry_insights": [
        "抗老需求年轻化，25-30岁成为抗老精华新增主力人群；",
        "成分党消费者决策更依赖社区测评内容，种草链路变长；",
        "医美术后修护场景带动高端修护类产品增长。",
    ],
    "sources": [
        {"title": "2026中国高端护肤市场洞察", "url": "https://example.com/report"},
        {"title": "美妆行业竞品投放策略分析", "url": "https://example.com/ads"},
    ],
}

MOCK_FORM = {
    "brand_name": "测试品牌 TestBrand",
    "product_name": "抗老精华 2026 新品上市",
    "ad_goal": "新品上市",
    "audience_desc": "25-35岁一二线城市女性，关注高端护肤，有一定消费力",
    "cooperation": "主推专利抗老成分，联合头部美妆达人种草",
    "competitors": "雅诗兰黛、兰蔻、赫莲娜",
    "budget": "300万",
    "regions": "全国一二线城市",
    "extra": "希望包含医美人群",
}


def main():
    # ---------- 1. 验证 old_skill 业务规则能被读到 ----------
    system_prompt, loaded = skill_rules.build_system_prompt()
    print("=" * 60)
    print("【1】old_skill 业务规则加载检查")
    print("  成功加载文件：", loaded)
    print("  system prompt 长度：{} 字".format(len(system_prompt)))
    assert len(loaded) == 3, "业务规则文件没有全部加载到！"
    assert "影像识别" in system_prompt and "种子人群包" in system_prompt

    user_prompt = skill_rules.build_user_prompt(MOCK_FORM)
    print("  user prompt 长度：{} 字".format(len(user_prompt)))

    # ---------- 2. 生成 Excel ----------
    print("=" * 60)
    print("【2】生成 Excel 测试")
    output_dir = history_store.ensure_output_dir()
    path = excel_builder.build_audience_excel(
        data=MOCK_DATA, output_dir=output_dir, form=MOCK_FORM,
        extra_meta={"web_search_used": True},
    )
    print("  生成路径：", path)
    assert os.path.exists(path), "Excel 文件没生成出来！"
    print("  文件大小：{} KB".format(round(os.path.getsize(path) / 1024, 1)))

    # ---------- 3. 回读校验内容 ----------
    print("=" * 60)
    print("【3】回读 Excel 校验结构")
    wb = openpyxl.load_workbook(path)
    print("  工作表：", wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    print("  主表尺寸：{} 行 x {} 列".format(ws.max_row, ws.max_column))
    print("  A1 标题：", ws["A1"].value)
    print("  合并单元格数量：", len(ws.merged_cells.ranges))

    # 检查四大板块标题都在
    all_text = []
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if v:
                all_text.append(str(v))
    joined = "\n".join(all_text)
    for must in ["数据说明：", "人群包汇总", "基础人群定向", "定制人群包圈选规则", "备注"]:
        assert must in joined, "缺少板块：{}".format(must)
        print("  ✓ 板块存在：", must)

    # 检查列宽
    widths = {c: ws.column_dimensions[c].width for c in ["A", "B", "C", "D"]}
    print("  列宽：", widths)
    assert widths["C"] == 68, "C 列宽度不对"

    # 检查人群包与子标签都写进去了
    for seg in MOCK_DATA["segments"]:
        assert seg["name"] in joined, "人群包缺失：{}".format(seg["name"])
        for t in seg["tags"]:
            assert t["logic"][:20] in joined, "圈选规则缺失：{}".format(t["tag"])
    print("  ✓ 全部 {} 个人群包 / {} 个子标签均已写入".format(
        len(MOCK_DATA["segments"]),
        sum(len(s["tags"]) for s in MOCK_DATA["segments"]),
    ))

    # 第二个工作表
    ws2 = wb["行业洞察与参考来源"]
    print("  洞察表尺寸：{} 行".format(ws2.max_row))
    assert ws2.max_row > 10

    # ---------- 4. 历史记录 + 配额 ----------
    print("=" * 60)
    print("【4】历史记录与配额测试")
    rec = history_store.add_history(
        brand_name=MOCK_FORM["brand_name"], product_name=MOCK_FORM["product_name"],
        ad_goal=MOCK_FORM["ad_goal"], requirement_summary=MOCK_FORM["audience_desc"],
        excel_path=path, segment_count=3, tag_count=10,
        web_search_used=True, full_form=MOCK_FORM,
    )
    print("  写入记录：", rec["timestamp"], rec["brand_name"])
    hist = history_store.load_history()
    print("  当前历史记录条数：", len(hist))
    assert len(hist) >= 1

    allowed, used, limit = history_store.check_quota(500)
    print("  配额：已用 {} / 上限 {}，允许调用：{}".format(used, limit, allowed))

    print("=" * 60)
    print("全部测试通过 ✓")
    return path


if __name__ == "__main__":
    main()
