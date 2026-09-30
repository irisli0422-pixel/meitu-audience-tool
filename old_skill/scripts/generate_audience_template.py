#!/usr/bin/env python3
"""
美图定制人群包 Excel 生成模板脚本
用法：修改 BRAND_NAME, BRAND_INFO, SEGMENTS 等变量后运行
依赖：openpyxl (pip install openpyxl)
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ============================================================
# 1. 品牌配置区（必改）
# ============================================================

BRAND_NAME = '品牌名称'  # 中文品牌名
BRAND_NAME_EN = 'BRAND'  # 英文品牌名
OUTPUT_FILENAME = f'{BRAND_NAME} x 美图定制人群包圈选规则.xlsx'
OUTPUT_DIR = '/Users/meitu/WorkBuddy/output'  # 输出目录

# 品牌配色（标题深色 / 板块中色 / 表头亮色）
COLOR_TITLE = '1A1A1A'    # 标题行背景
COLOR_SECTION = '2E5C8A'  # 板块标题背景
COLOR_HEADER = '3B7DD8'   # 表头背景
COLOR_BASIC = 'E8F0FE'    # 基础定向背景

# 数据说明
DATA_NOTES = [
    '1. 整体逻辑通过XXX消费人群特征匹配美图人群画像；',
    '2. 涉及的APP：美图系场景（美图秀秀、美颜相机等）；',
    '3. 人群包采用近6个月中符合圈选规则的人群；',
    '4. 品牌背景信息...',  # ← 替换为品牌信息
    '5. 本次核心场景：XXX、XXX、XXX。',
]

# 人群包汇总公式
SUMMARY_TEXT = '【基础定向 ∪ 影像识别XXX人群 ∪ 场景识别XXX人群】∩【人群包A ∪ 人群包B ∪ ... ∪ 种子人群包】'

# 基础人群定向
BASIC_TARGETING = [
    ('', '性别维度', '女 / 男（说明偏向原因），整体男女全覆盖', ''),
    ('', '年龄维度', '18-40岁（核心人群说明）', ''),
    ('', '兴趣爱好维度', '安装应用偏好：XXX\n社区兴趣偏好：XXX', ''),
    ('', '中高消费维度', '中等及以上消费力（月收入XXXX+），对XXX有消费意愿', ''),
]

# ============================================================
# 2. 人群包数据区（必改）
# ============================================================
# 格式：(人群包名称, 子标签名称, 圈选规则, 预估量级)

SEGMENTS = [
    # 1. 本竞品人群包（必选）
    ('本竞品人群包', '广告点击行为-品牌',
     '在美图系平台的广告点击行为数据判断出用户对相关广告感兴趣，广告点击兴趣=品牌A/品牌B/品牌C，频次=180天', '120w'),
    ('本竞品人群包', '影像识别-竞品品牌',
     '近180天内通过影像识别侦察到【竞品A、竞品B、竞品C】品牌形象的用户', '280w'),

    # 2-N. 品牌定制人群包（根据品牌特性设计）
    # ... 添加更多人群包 ...

    # 最后. 种子人群包（必选）
    ('种子人群包', '历史互动人群',
     '2024-2025年度在美图秀秀APP点击过XX品牌广告、参与过XX话题互动、使用过XX联名配方/贴纸/滤镜的人群（去重种子人群包）', '85w'),
    ('种子人群包', '品牌搜索人群',
     '近365天内在美图系平台搜索过XX/XX/XX等关键词的用户', '52w'),
]

# 各人群包行背景色（浅色交替）
SEGMENT_COLORS = {
    '本竞品人群包': 'FFF3E0',
    # ... 每个人群包一个浅色 ...
    '种子人群包': 'EDEDED',
}

# ============================================================
# 3. 备注区
# ============================================================
REMARKS = [
    '1. 以上人群包量级为预估量级，实际投放时以系统最终圈选结果为准；',
    '2. 各人群包之间采用"并集（∪）"逻辑组合，基础定向与定制人群包之间采用"交集（∩）"逻辑；',
    '3. 影像识别基于美图AI图像识别技术，识别用户上传/编辑图片中的品牌LOGO、场景、物品等元素；',
    '4. LBS场景识别基于用户打开APP时的地理位置信息，匹配线下消费场所POI数据；',
    '5. 原生需求指用户在美图秀秀产品中主动使用的功能/配方/滤镜等行为数据；',
    '6. 品牌核心信息...',  # ← 替换
    '7. 核心人群优先级：XXX人群包(高转化) → XXX人群包(高复购) → XXX人群包(高传播)。',
]

# ============================================================
# 4. 生成逻辑（通常无需修改）
# ============================================================

def generate():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{BRAND_NAME_EN}定制人群包"

    # 样式定义
    title_font = Font(name='微软雅黑', size=16, bold=True, color='FFFFFF')
    section_font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    content_font = Font(name='微软雅黑', size=10, color='333333')
    content_bold = Font(name='微软雅黑', size=10, bold=True, color='333333')
    note_font = Font(name='微软雅黑', size=9, color='666666')

    title_fill = PatternFill(start_color=COLOR_TITLE, end_color=COLOR_TITLE, fill_type='solid')
    section_fill = PatternFill(start_color=COLOR_SECTION, end_color=COLOR_SECTION, fill_type='solid')
    header_fill = PatternFill(start_color=COLOR_HEADER, end_color=COLOR_HEADER, fill_type='solid')
    basic_fill = PatternFill(start_color=COLOR_BASIC, end_color=COLOR_BASIC, fill_type='solid')
    white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )

    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    left_top_align = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # 列宽
    for col, width in {'A': 20, 'B': 26, 'C': 68, 'D': 14}.items():
        ws.column_dimensions[col].width = width

    row = 1

    # 标题行
    ws.merge_cells(f'A{row}:D{row}')
    cell = ws.cell(row=row, column=1, value=f'{BRAND_NAME} x 美图定制人群包圈选规则')
    cell.font = title_font
    cell.fill = title_fill
    cell.alignment = center_align
    ws.row_dimensions[row].height = 40
    row += 2

    # 数据说明
    ws.merge_cells(f'A{row}:D{row}')
    cell = ws.cell(row=row, column=1, value='数据说明：')
    cell.font = section_font
    cell.fill = section_fill
    cell.alignment = left_align
    ws.row_dimensions[row].height = 28
    row += 1

    for note in DATA_NOTES:
        ws.merge_cells(f'A{row}:D{row}')
        cell = ws.cell(row=row, column=1, value=note)
        cell.font = note_font
        cell.alignment = left_align
        ws.row_dimensions[row].height = 20
        row += 1
    row += 1

    # 人群包汇总
    ws.merge_cells(f'A{row}:D{row}')
    cell = ws.cell(row=row, column=1, value='人群包汇总')
    cell.font = section_font
    cell.fill = section_fill
    cell.alignment = left_align
    ws.row_dimensions[row].height = 28
    row += 1

    ws.merge_cells(f'A{row}:D{row}')
    cell = ws.cell(row=row, column=1, value=SUMMARY_TEXT)
    cell.font = content_font
    cell.fill = basic_fill
    cell.alignment = left_align
    ws.row_dimensions[row].height = 36
    row += 2

    # 基础人群定向
    ws.merge_cells(f'A{row}:D{row}')
    cell = ws.cell(row=row, column=1, value='基础人群定向')
    cell.font = section_font
    cell.fill = section_fill
    cell.alignment = left_align
    ws.row_dimensions[row].height = 28
    row += 1

    headers = ['基础人群定向', '人群维度', '秀秀单标签逻辑', '人群包量级']
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[row].height = 28
    row += 1

    basic_start = row
    for i, (col_a, col_b, col_c, col_d) in enumerate(BASIC_TARGETING):
        fill = basic_fill if i % 2 == 0 else white_fill
        for c_idx, val in enumerate([col_a, col_b, col_c, col_d], 1):
            cell = ws.cell(row=row, column=c_idx, value=val)
            cell.fill = fill
            cell.font = content_font if c_idx != 2 else content_bold
            cell.alignment = left_align if c_idx == 3 else center_align
            cell.border = thin_border
        ws.row_dimensions[row].height = 50 if i == 2 else 24
        row += 1

    ws.merge_cells(f'A{basic_start}:A{row-1}')
    merged = ws.cell(row=basic_start, column=1, value='基础人群定向')
    merged.font = content_bold
    merged.alignment = center_align
    merged.fill = basic_fill
    row += 1

    # 定制人群包圈选规则
    ws.merge_cells(f'A{row}:D{row}')
    cell = ws.cell(row=row, column=1, value=f'{BRAND_NAME}定制人群包圈选规则')
    cell.font = section_font
    cell.fill = section_fill
    cell.alignment = left_align
    ws.row_dimensions[row].height = 28
    row += 1

    headers2 = ['人群包', '人群标签', '秀秀单标签逻辑之一', '人群包量级']
    for col_idx, h in enumerate(headers2, 1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[row].height = 28
    row += 1

    # 写入人群包数据
    current_segment = None
    segment_start_row = row

    for seg_name, tag, logic, scale in SEGMENTS:
        fill_color = SEGMENT_COLORS.get(seg_name, 'FFFFFF')
        fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')

        if seg_name != current_segment:
            if current_segment is not None and row - 1 > segment_start_row:
                ws.merge_cells(f'A{segment_start_row}:A{row-1}')
            segment_start_row = row
            current_segment = seg_name
            ws.cell(row=row, column=1, value=seg_name)
        else:
            ws.cell(row=row, column=1, value='')

        ws.cell(row=row, column=2, value=tag)
        ws.cell(row=row, column=3, value=logic)
        ws.cell(row=row, column=4, value=scale)

        for c in range(1, 5):
            cell = ws.cell(row=row, column=c)
            cell.font = content_font if c != 1 else content_bold
            cell.alignment = center_align if c in (1, 2, 4) else left_top_align
            cell.border = thin_border
            cell.fill = fill

        logic_lines = logic.count('\n') + 1
        base_height = max(36, min(90, len(logic) // 2 + logic_lines * 18))
        ws.row_dimensions[row].height = base_height
        row += 1

    if row - 1 > segment_start_row:
        ws.merge_cells(f'A{segment_start_row}:A{row-1}')
    row += 1

    # 备注
    ws.merge_cells(f'A{row}:D{row}')
    cell = ws.cell(row=row, column=1, value='备注')
    cell.font = section_font
    cell.fill = section_fill
    cell.alignment = left_align
    ws.row_dimensions[row].height = 28
    row += 1

    for remark in REMARKS:
        ws.merge_cells(f'A{row}:D{row}')
        cell = ws.cell(row=row, column=1, value=remark)
        cell.font = note_font
        cell.alignment = left_align
        ws.row_dimensions[row].height = 20
        row += 1

    # 保存
    import os
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output_path = os.path.join(OUTPUT_DIR, OUTPUT_FILENAME)
    wb.save(output_path)
    print(f'文件已保存至: {output_path}')
    return output_path


if __name__ == '__main__':
    generate()
